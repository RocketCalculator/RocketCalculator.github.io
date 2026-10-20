"""Dump every converted workbook (derived/xlsx) to data/spreadsheets/<name>.json:
sheets -> cells with cached value + formula, so a web port can re-implement the maths.
Also dumps VBA macros (source/.../media/*.xls) to data/vba/<name>.bas via olevba.
Run with the venv python that has openpyxl + oletools."""
import json, os, sys, glob, subprocess
from openpyxl import load_workbook
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
out = f'{ROOT}/data/spreadsheets'; os.makedirs(out, exist_ok=True)
vout = f'{ROOT}/data/vba'; os.makedirs(vout, exist_ok=True)
srcs = glob.glob(f'{ROOT}/derived/xlsx/*.xlsx') + glob.glob(f'{ROOT}/source/rsandt/media/*.xlsx')
for p in sorted(srcs):
    name = os.path.splitext(os.path.basename(p))[0]
    try:
        wf = load_workbook(p, data_only=False); wv = load_workbook(p, data_only=True)
    except Exception as e:
        print('FAIL', name, e); continue
    book = {}
    for ws in wf.worksheets:
        vs = wv[ws.title]; cells = {}
        for row in ws.iter_rows():
            for c in row:
                if c.value is None: continue
                d = {'v': vs[c.coordinate].value}
                if isinstance(c.value, str) and c.value.startswith('='): d['f'] = c.value
                else: d['v'] = c.value
                cells[c.coordinate] = d
        book[ws.title] = {'dims': ws.dimensions, 'cells': cells}
    json.dump(book, open(f'{out}/{name}.json', 'w'), default=str)
    print('ok', name, {k: len(v['cells']) for k, v in book.items()})
for p in sorted(glob.glob(f'{ROOT}/source/rsandt/media/*.xls')):
    name = os.path.splitext(os.path.basename(p))[0]
    r = subprocess.run([sys.executable, '-m', 'oletools.olevba', '-c', '--decode', p], capture_output=True, text=True)
    if 'VBA MACRO' in r.stdout:
        open(f'{vout}/{name}.bas', 'w').write(r.stdout)
        print('vba', name)
